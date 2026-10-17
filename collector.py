import os
import sys
import io
import json
import textwrap
import subprocess
from datetime import datetime
from typing import Dict, List, Optional, Tuple, Union

import pandas as pd

# Streamlit Cloud에서 requirement.txt(단수형) 파일명만 업로드된 경우
# 패키지가 자동 설치되지 않아 발생하는 ModuleNotFoundError 방지용 자동 설치 로직
try:
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from dotenv import load_dotenv
    from google import genai
    from google.genai import types
except ModuleNotFoundError:
    subprocess.check_call(
        [
            sys.executable,
            "-m",
            "pip",
            "install",
            "openpyxl>=3.1.2",
            "google-genai>=1.0.0",
            "python-dotenv>=1.0.0",
        ]
    )
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from dotenv import load_dotenv
    from google import genai
    from google.genai import types

load_dotenv()

# 기본 파일 경로 설정
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_EXCEL_PATH = os.path.join(BASE_DIR, "Gemini_English.xlsx")
USERS_FILE_PATH = os.path.join(BASE_DIR, "users.json")

# 기본 관리자 계정 정보 (작업지도서 명세)
ADMIN_ID = "admin"
ADMIN_DEFAULT_PW = "jhlee_english01"

# 출력/저장 데이터 표준 컬럼 포맷 ("User" 첫번째 열 추가)
COLUMNS = [
    "User",
    "Topic",
    "Date",
    "Who",
    "Conversation History",
    "Error Correction",
]


# ==========================================
# 1. 회원 관리 (신규등록 / 로그인 / 초기화)
# ==========================================
def init_users_db() -> Dict[str, str]:
    """
    회원 정보를 저장하는 users.json 파일을 초기화하고 불러옵니다.
    기본 관리자 계정(ID: admin, PW: jhlee_english01)을 항상 보장합니다.
    """
    users: Dict[str, str] = {}
    if os.path.exists(USERS_FILE_PATH):
        try:
            with open(USERS_FILE_PATH, "r", encoding="utf-8") as f:
                loaded = json.load(f)
                if isinstance(loaded, dict):
                    users = loaded
        except Exception:
            users = {}

    if ADMIN_ID not in users:
        users[ADMIN_ID] = ADMIN_DEFAULT_PW
        save_users_db(users)

    return users


def save_users_db(users: Dict[str, str]) -> None:
    """회원 정보를 users.json 파일에 저장합니다."""
    with open(USERS_FILE_PATH, "w", encoding="utf-8") as f:
        json.dump(users, f, ensure_ascii=False, indent=2)


def is_user_registered(user_id: str) -> bool:
    """해당 ID가 이미 등록되어 있는지 확인합니다."""
    users = init_users_db()
    return user_id.strip() in users


def register_user(user_id: str, pw1: str, pw2: str) -> Tuple[bool, str]:
    """
    신규 회원을 등록합니다.
    - 이미 등록된 회원의 경우: "회원 등록이 이미 되었습니다" 메시지 반환
    - 비밀번호 2회 입력 불일치 시 오류 메시지 반환
    """
    uid = user_id.strip()
    if not uid:
        return False, "ID를 입력해 주세요."
    if not pw1 or not pw2:
        return False, "Password를 2회 모두 입력해 주세요."

    users = init_users_db()
    if uid in users:
        return False, "회원 등록이 이미 되었습니다"

    if pw1 != pw2:
        return False, "입력하신 두 Password가 일치하지 않습니다."

    users[uid] = pw1
    save_users_db(users)
    return True, f"'{uid}' 신규 회원 등록이 완료되었습니다. 프로그램 사용 권한이 부여되었습니다."


def verify_login(user_id: str, password: str) -> Tuple[bool, str]:
    """
    Log In 시 ID와 Password가 일치하는지 검증합니다.
    """
    uid = user_id.strip()
    if not uid or not password:
        return False, "ID와 Password를 모두 입력해 주세요."

    users = init_users_db()
    if uid not in users or users[uid] != password:
        return False, "입력된 ID나 Password가 일치하지 않습니다. 다시 확인해 주세요."

    return True, f"환영합니다, '{uid}'님! 정상적으로 로그인되었습니다."


def reset_user_account(
    admin_id: str,
    admin_pw: str,
    target_user_id: str,
    new_pw1: str = "",
    new_pw2: str = "",
    delete_account: bool = False,
) -> Tuple[bool, str]:
    """
    관리자 계정(ID: admin, PW: jhlee_english01)을 통해
    분실된 회원의 비밀번호를 재설정하거나 회원 계정을 초기화(삭제)합니다.
    """
    users = init_users_db()
    aid = admin_id.strip()
    tid = target_user_id.strip()

    expected_admin_pw = users.get(ADMIN_ID, ADMIN_DEFAULT_PW)
    if aid != ADMIN_ID or admin_pw != expected_admin_pw:
        return False, "관리자 ID(admin) 또는 관리자 Password가 일치하지 않습니다."

    if not tid:
        return False, "초기화할 대상 회원 ID를 입력해 주세요."

    if tid not in users:
        return False, f"'{tid}' 회원은 등록되어 있지 않습니다."

    if delete_account:
        if tid == ADMIN_ID:
            users[ADMIN_ID] = ADMIN_DEFAULT_PW
            save_users_db(users)
            return True, "관리자(admin) 비밀번호가 기본값(jhlee_english01)으로 초기화되었습니다."
        del users[tid]
        save_users_db(users)
        return True, f"'{tid}' 회원 계정이 초기화(삭제)되었습니다. 신규등록에서 다시 등록하실 수 있습니다."

    if not new_pw1 or not new_pw2:
        return False, "새로 설정할 Password를 2회 모두 입력해 주세요."
    if new_pw1 != new_pw2:
        return False, "새로 입력한 두 Password가 일치하지 않습니다."

    users[tid] = new_pw1
    save_users_db(users)
    return True, f"'{tid}' 회원의 Password가 성공적으로 초기화(변경)되었습니다."


def get_registered_user_ids(admin_id: str, admin_pw: str) -> Tuple[bool, List[str], str]:
    """관리자 인증 후 등록된 전체 회원 ID 목록을 조회합니다 (ID 분실 시 확인 지원)."""
    users = init_users_db()
    expected_admin_pw = users.get(ADMIN_ID, ADMIN_DEFAULT_PW)
    if admin_id.strip() != ADMIN_ID or admin_pw != expected_admin_pw:
        return False, [], "관리자 ID(admin) 또는 관리자 Password가 일치하지 않습니다."
    return True, list(users.keys()), "등록된 회원 ID 목록 조회가 완료되었습니다."


# ==========================================
# 2. Gemini API 통신 (01_Gemini_API_Eng_Study.py 동일 로직)
# ==========================================
def get_api_key(custom_api_key: Optional[str] = None) -> Optional[str]:
    """환경변수(.env), Streamlit secrets, 또는 사용자 입력에서 GOOGLE_API_KEY를 가져옵니다."""
    if custom_api_key and custom_api_key.strip():
        return custom_api_key.strip()

    env_key = os.getenv("GOOGLE_API_KEY")
    if env_key and env_key.strip():
        return env_key.strip()

    try:
        import streamlit as st
        if "GOOGLE_API_KEY" in st.secrets:
            return str(st.secrets["GOOGLE_API_KEY"]).strip()
    except Exception:
        pass

    return None


def create_gemini_chat_session(
    api_key: Optional[str] = None,
    model_name: str = "gemini-2.5-pro",
    history: Optional[ list ] = None,
):
    """
    01_Gemini_API_Eng_Study.py와 동일하게 신규 Gemini SDK Client 및 Chat 세션을 생성합니다.
    주의: google-genai SDK의 Client 객체는 참조가 사라져 가비지 컬렉션(GC)될 때
    __del__()에서 내부 httpx 클라이언트를 close()하므로, chat_session 내부에
    client 강한 참조(_keep_alive_client)를 함께 보관하여 연결 종료를 방지합니다.
    """
    resolved_key = get_api_key(api_key)
    if not resolved_key:
        raise ValueError("GOOGLE_API_KEY가 설정되어 있지 않습니다. .env 파일 또는 화면 좌측 설정에 API Key를 입력해 주세요.")

    client = genai.Client(api_key=resolved_key)
    chat_kwargs = {
        "model": model_name,
        "config": types.GenerateContentConfig(
            response_mime_type="application/json"
        ),
    }
    if history:
        chat_kwargs["history"] = history

    chat_session = client.chats.create(**chat_kwargs)
    # Client 객체가 GC에 의해 소멸되어 httpx client가 닫히는 현상 원천 차단
    setattr(chat_session, "_keep_alive_client", client)
    return client, chat_session


def ask_gemini_tutor(
    chat_session,
    topic: str,
    user_message: str,
    api_key: Optional[str] = None,
    model_name: str = "gemini-2.5-pro",
) -> Tuple[Dict[str, str], object]:
    """
    01_Gemini_API_Eng_Study.py의 GeminiThread.run()과 동일한 프롬프트 및 JSON 응답 구조를 수행합니다.
    만약 네트워크 유휴 등으로 인해 client가 닫힌 경우(Cannot send a request, as the client has been closed),
    기존 대화 히스토리를 유지한 채 새 세션을 재생성하여 자동 재시도합니다.
    """
    prompt = (
        f"Current Conversation Topic: {topic}\n"
        f"User Message: {user_message}\n\n"
        "Act as an English conversation tutor. Respond strictly in JSON format with exactly two keys:\n"
        "1. 'correction': The corrected version of the user's sentence (Best Expression). If the user's sentence is grammatically correct and natural, leave it as an empty string \"\".\n"
        "2. 'response': Your natural conversational reply to the user's message."
    )

    try:
        response = chat_session.send_message(prompt)
    except RuntimeError as e:
        if "client has been closed" in str(e).lower():
            prev_history = None
            try:
                prev_history = chat_session.get_history()
            except Exception:
                prev_history = None
            _, chat_session = create_gemini_chat_session(
                api_key=api_key,
                model_name=model_name,
                history=prev_history,
            )
            response = chat_session.send_message(prompt)
        else:
            raise

    raw_text = (response.text or "").strip()

    # 마크다운 코드블록이 포함된 경우 제거 방어 로직
    if raw_text.startswith("```"):
        lines = raw_text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].startswith("```"):
            lines = lines[:-1]
        raw_text = "\n".join(lines).strip()

    result = json.loads(raw_text)
    if isinstance(result, list) and len(result) > 0 and isinstance(result[0], dict):
        result = result[0]

    correction = str(result.get("correction", "") or "").strip()
    gemini_reply = str(result.get("response", "") or "").strip()
    return {
        "correction": correction,
        "response": gemini_reply,
    }, chat_session


# ==========================================
# 3. 영어회화 로그 수집 및 엑셀 관리 (Gemini_English.xlsx)
# ==========================================
def _normalize_dataframe(df: Optional[pd.DataFrame]) -> pd.DataFrame:
    """
    기존 Gemini_English.xlsx에 'User' 열이 없는 경우 첫번째 열에 'User'를 추가하고,
    표준 컬럼 순서 ['User', 'Topic', 'Date', 'Who', 'Conversation History', 'Error Correction']로 정렬합니다.
    """
    if df is None or df.empty:
        return pd.DataFrame(columns=COLUMNS)

    df = df.copy()
    if "User" not in df.columns:
        df.insert(0, "User", "")

    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""

    df = df[COLUMNS].fillna("")
    for col in COLUMNS:
        df[col] = df[col].astype(str).replace({"nan": "", "None": ""})
    return df


def format_cell_linebreaks(text: str, width: int = 55) -> str:
    """
    문장 길이가 엑셀 셀 크기를 초과할 경우 셀 안에 글자가 잘 표시되도록
    줄바꿈(\n, 엑셀 셀 내 줄바꿈 Shift+Enter/Alt+Enter)을 정리합니다.
    """
    if not text:
        return ""
    lines = str(text).splitlines()
    wrapped_lines = []
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
        if len(line_str) <= width:
            wrapped_lines.append(line_str)
        else:
            wrapped_lines.extend(
                textwrap.wrap(
                    line_str,
                    width=width,
                    break_long_words=False,
                    break_on_hyphens=False,
                )
            )
    return "\n".join(wrapped_lines)


def write_formatted_excel(
    df: pd.DataFrame,
    target_or_buffer: Union[str, io.BytesIO],
    wrap_width_history: int = 55,
    wrap_width_correction: int = 45,
) -> None:
    """
    DataFrame을 Gemini_English.xlsx 포맷(첫번째 열 User 포함)으로 저장하고,
    openpyxl을 사용하여 셀 내부 줄바꿈(wrap_text=True) 및 열 너비를 자동 맞춤합니다.
    """
    norm_df = _normalize_dataframe(df).copy()

    # 셀 폭을 초과하는 긴 문장에 대해 줄바꿈(\n) 정리
    norm_df["Conversation History"] = norm_df["Conversation History"].apply(
        lambda x: format_cell_linebreaks(x, width=wrap_width_history)
    )
    norm_df["Error Correction"] = norm_df["Error Correction"].apply(
        lambda x: format_cell_linebreaks(x, width=wrap_width_correction)
    )

    if isinstance(target_or_buffer, str):
        save_dir = os.path.dirname(target_or_buffer)
        if save_dir:
            os.makedirs(save_dir, exist_ok=True)

    with pd.ExcelWriter(target_or_buffer, engine="openpyxl") as writer:
        norm_df.to_excel(writer, index=False, sheet_name="Log")
        worksheet = writer.sheets["Log"]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # 제목 행 고정 (엑셀에서 스크롤 시 1행 고정)
        worksheet.freeze_panes = "A2"

        # 전체 셀 서식 (줄바꿈 wrap_text=True 및 상단 정렬)
        for row_idx, row in enumerate(
            worksheet.iter_rows(min_row=1, max_col=len(COLUMNS), max_row=worksheet.max_row),
            start=1,
        ):
            max_lines_in_row = 1
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.value:
                        line_count = str(cell.value).count("\n") + 1
                        if line_count > max_lines_in_row:
                            max_lines_in_row = line_count

            # 줄바꿈된 줄 수에 맞춰 행 높이 자동 조정
            if row_idx == 1:
                worksheet.row_dimensions[row_idx].height = 24
            else:
                worksheet.row_dimensions[row_idx].height = max(20, max_lines_in_row * 18)

        # 컬럼 너비 설정 (첫번째 열 User 추가 반영)
        worksheet.column_dimensions["A"].width = 15  # User
        worksheet.column_dimensions["B"].width = 16  # Topic
        worksheet.column_dimensions["C"].width = 21  # Date
        worksheet.column_dimensions["D"].width = 12  # Who
        worksheet.column_dimensions["E"].width = 55  # Conversation History
        worksheet.column_dimensions["F"].width = 45  # Error Correction


def load_all_logs(excel_path: str = DEFAULT_EXCEL_PATH) -> pd.DataFrame:
    """전체 영어회화 로그 데이터를 불러옵니다 (User 컬럼 자동 보장)."""
    if not os.path.exists(excel_path):
        return pd.DataFrame(columns=COLUMNS)
    try:
        df = pd.read_excel(excel_path)
        return _normalize_dataframe(df)
    except Exception:
        return pd.DataFrame(columns=COLUMNS)


def load_user_logs(user_id: str, excel_path: str = DEFAULT_EXCEL_PATH) -> pd.DataFrame:
    """
    로그인한 사용자(user_id) 본인의 영어회화 로그만 필터링하여 반환합니다.
    다른 사용자의 ID에 해당하는 로그는 절대 포함하지 않습니다.
    """
    uid = user_id.strip()
    all_df = load_all_logs(excel_path)
    if all_df.empty or not uid:
        return pd.DataFrame(columns=COLUMNS)

    user_df = all_df[all_df["User"].astype(str).str.strip() == uid].copy()
    user_df.reset_index(drop=True, inplace=True)
    return user_df


def append_conversation_turn(
    user_id: str,
    topic: str,
    user_text: str,
    correction: str,
    gemini_response: str,
    user_time: Optional[str] = None,
    gemini_time: Optional[str] = None,
    excel_path: str = DEFAULT_EXCEL_PATH,
) -> pd.DataFrame:
    """
    한 턴의 영어회화(Me 입력 + Error Correction + Gemini 응답)를
    Gemini_English.xlsx 포맷(첫번째 열 User 포함)으로 즉시 저장합니다.
    """
    uid = user_id.strip()
    t_user = user_time or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    t_gemini = gemini_time or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    new_rows = [
        {
            "User": uid,
            "Topic": topic.strip(),
            "Date": t_user,
            "Who": "Me",
            "Conversation History": user_text.strip(),
            "Error Correction": correction.strip() if correction else "",
        },
        {
            "User": uid,
            "Topic": topic.strip(),
            "Date": t_gemini,
            "Who": "Gemini",
            "Conversation History": gemini_response.strip(),
            "Error Correction": "",
        },
    ]

    existing_df = load_all_logs(excel_path)
    new_df = pd.DataFrame(new_rows, columns=COLUMNS)
    combined_df = pd.concat([existing_df, new_df], ignore_index=True)
    write_formatted_excel(combined_df, excel_path)
    return load_user_logs(uid, excel_path)


def export_user_logs_to_bytes(user_id: str, excel_path: str = DEFAULT_EXCEL_PATH) -> bytes:
    """
    로그인한 ID에 해당하는 영어회화 로그만 추출하여
    줄바꿈 서식이 적용된 엑셀 바이트(BytesIO)로 반환합니다.
    """
    user_df = load_user_logs(user_id, excel_path)
    buffer = io.BytesIO()
    write_formatted_excel(user_df, buffer)
    buffer.seek(0)
    return buffer.getvalue()


def save_user_logs_to_path(
    user_id: str,
    target_file_path: str,
    excel_path: str = DEFAULT_EXCEL_PATH,
) -> Tuple[bool, str]:
    """
    로그인한 ID에 해당하는 영어회화 로그 데이터를 사용자가 지정한 폴더와 파일명으로 저장합니다.
    """
    try:
        user_df = load_user_logs(user_id, excel_path)
        if not target_file_path.lower().endswith(".xlsx"):
            target_file_path += ".xlsx"
        write_formatted_excel(user_df, target_file_path)
        return True, f"엑셀 파일이 성공적으로 저장되었습니다: {target_file_path}"
    except Exception as e:
        return False, f"엑셀 파일 저장 중 오류가 발생했습니다: {str(e)}"


def delete_user_logs(user_id: str, excel_path: str = DEFAULT_EXCEL_PATH) -> Tuple[bool, int, str]:
    """
    로그인한 ID(user_id)에 해당하는 모든 영어회화 로그 데이터를 삭제합니다.
    다른 사용자의 데이터는 안전하게 유지합니다.
    """
    uid = user_id.strip()
    if not uid:
        return False, 0, "로그인된 사용자 ID가 없습니다."

    all_df = load_all_logs(excel_path)
    if all_df.empty:
        return True, 0, "삭제할 영어회화 로그 데이터가 없습니다."

    mask_user = all_df["User"].astype(str).str.strip() == uid
    deleted_count = int(mask_user.sum())

    remaining_df = all_df[~mask_user].copy()
    remaining_df.reset_index(drop=True, inplace=True)
    write_formatted_excel(remaining_df, excel_path)

    return True, deleted_count, f"'{uid}'님의 영어회화 로그 데이터 {deleted_count}건이 모두 삭제되었습니다."
